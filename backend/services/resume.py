from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import pymupdf
from docx import Document
from fastapi import UploadFile

from backend.intelligence.gateway import ModelGateway
from backend.schemas.domain import (
    CandidateProfileData,
    PersonalProfileData,
    ResumeData,
    ResumeImportData,
)

ALLOWED_SUFFIXES = {".pdf", ".docx", ".txt"}


def extract_local_text(path: Path, suffix: str) -> str:
    if suffix == ".pdf":
        with pymupdf.open(path) as document:
            text = "\n".join(page.get_text() for page in document)
    elif suffix == ".docx":
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(paragraphs + table_rows)
    else:
        text = path.read_text(encoding="utf-8")
    return text.strip()


def _mock_import(text: str, filename: str) -> ResumeImportData:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = next((line for line in lines[1:] if len(line.split()) <= 8), None)
    skill_words = ("Python", "FastAPI", "SQL", "Django", "JavaScript", "Docker", "Git")
    skills = [skill for skill in skill_words if skill.lower() in text.lower()]
    return ResumeImportData(
        profile=PersonalProfileData(full_name=lines[0] if lines else None),
        resume=ResumeData(
            name=Path(filename).stem or "Резюме",
            desired_title=title,
            skills=skills,
            about=" ".join(lines[1:3]),
            original_filename=filename,
        ),
    )


async def save_and_extract(file: UploadFile, gateway: ModelGateway) -> tuple[Path, ResumeImportData]:
    """Persist an import only after local extraction and structured parsing succeed."""
    filename = file.filename or "resume"
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError("Поддерживаются только PDF, DOCX и TXT")

    folder = Path("data/resumes")
    folder.mkdir(parents=True, exist_ok=True)
    path = (folder / f"{uuid4().hex}{suffix}").resolve()
    try:
        path.write_bytes(await file.read())
        text = extract_local_text(path, suffix)
        if not text:
            if suffix == ".pdf":
                raise ValueError("PDF не содержит текстового слоя; OCR в MVP не поддерживается")
            raise ValueError("Резюме не содержит текста")

        if getattr(gateway, "provider", None) == "mock":
            parsed = _mock_import(text, filename)
        else:
            parsed = await gateway.structured(
                "profile",
                {
                    "resume_text": text[:30_000],
                    "output_contract": "Return only profile and resume fields. Do not return fact records or source excerpts.",
                },
                ResumeImportData,
            )
        parsed.resume.original_filename = filename
        parsed.resume.original_path = str(path)
        return path, parsed
    except Exception:
        path.unlink(missing_ok=True)
        raise


def profile_from_import(data: ResumeImportData) -> CandidateProfileData:
    return CandidateProfileData.model_validate(data.profile.model_dump())


def resume_from_import(data: ResumeImportData, *, path: Path, filename: str) -> ResumeData:
    values = data.resume.model_dump()
    values["original_filename"] = filename
    values["original_path"] = str(path)
    return ResumeData.model_validate(values)
