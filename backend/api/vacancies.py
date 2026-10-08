"""SQL-backed vacancy list, detail and export endpoints.

The legacy API materialises every vacancy and evaluation before filtering.  The
router in this module deliberately selects only scalar projection columns for
lists and lets the database do filtering, counting, ordering and pagination.
It is the sole owner of the vacancy HTTP routes.
"""

from __future__ import annotations

import csv
import io
import json
import tempfile
import xml.etree.ElementTree as ET
from collections.abc import Iterator
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from sqlalchemy import Integer, String, case, cast, exists, func, literal, or_, select
from sqlalchemy.orm import Session, sessionmaker

from backend.persistence.database import get_db
from backend.persistence.models import Application, Evaluation, Vacancy

router = APIRouter(prefix="/api")

VACANCY_SCORE_KEYS = (
    "tasks", "skills", "experience_depth", "role_match", "industry", "special_requirements"
)
VACANCY_SCORE_ALIASES = {
    "experience_depth": "required_years", "role_match": "title", "special_requirements": "languages"
}
VacancySort = Literal[
    "id", "title", "state", "date", "site", "total_score", *VACANCY_SCORE_KEYS
]
VacancySortDirection = Literal["asc", "desc"]
VacancyExportFormat = Literal["csv", "xlsx", "xml"]
VacancyStatusGroup = Literal["SUCCESS", "PARTIAL", "PROCESSING", "REJECTED", "ERROR", "CANCELLED"]
VACANCY_STATUS_GROUPS = {
    "SUCCESS": frozenset({"SUBMITTED", "ALREADY_APPLIED", "REPORTED"}),
    "PARTIAL": frozenset({"PARTIAL"}),
    "PROCESSING": frozenset({"EXTRACTED", "EVALUATING", "READY_TO_SUBMIT", "READY_TO_REPORT", "SUBMITTING"}),
    "REJECTED": frozenset({"REJECTED_BY_MODEL"}),
    "ERROR": frozenset({"ERROR", "UNCONFIRMED", "SUBMISSION_UNCONFIRMED"}),
    "CANCELLED": frozenset({"CANCELLED"}),
}
VACANCY_SCORE_COLUMNS = {"total_score": Evaluation.total_score, **{key: getattr(Evaluation, key) for key in VACANCY_SCORE_KEYS}}
VACANCY_EXPORT_HEADERS = (
    "Номер вакансии", "Название вакансии", "Компания", "Сайт", "Дата", "Общий балл",
    "Задачи", "Навыки", "Опыт", "Роль", "Сфера", "Особые требования"
)
_EXPORT_BATCH_SIZE = 500


def _public_status(state: str | None) -> str:
    return "ERROR" if state == "UNCONFIRMED" else (state or "")


def _status_time(value: datetime | None, site: str | None = None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.isoformat(timespec="seconds")


def _export_status_time(value: datetime | None, site: str | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    value = value.astimezone()
    return value.strftime("%d.%m.%Y %H:%M")


def _error(row) -> tuple[str, str]:
    legacy = row.state in {"UNCONFIRMED", "SUBMISSION_UNCONFIRMED"} or getattr(row, "effective_state", None) == "SUBMISSION_UNCONFIRMED"
    message = row.error_message
    if not isinstance(message, str) or not message.strip() or message.strip() == "[удалено]":
        message = None
    return (
        row.error_code or ("SUBMISSION_UNCONFIRMED" if legacy else "VACANCY_PROCESSING_FAILED"),
        message or (
            "Не удалось подтвердить отправку отклика после нескольких попыток"
            if legacy else "Вакансия не обработана из-за ошибки"
        ),
    )


def _evaluation_summary(row, data: dict | None = None) -> dict | None:
    if getattr(row, "history_origin_id", None) is not None or getattr(row, "history_invalid", False):
        return None
    if row.evaluation_id is None:
        return None
    result = {
        "score": row.total_score, "decision": row.decision,
        "confidence": row.confidence, "category": row.category,
        "score_breakdown": {key: getattr(row, key) for key in VACANCY_SCORE_KEYS if getattr(row, key) is not None},
    }
    if data is not None:
        result.update({key: value for key, value in data.items() if key != "flag_matches"})
        result["score"] = row.total_score
        result["score_breakdown"] = data.get("score_breakdown", result["score_breakdown"])
    return result


def _history_projection():
    """Bounded SQL projection shared by list, detail, count and export paths."""
    seed_root = Vacancy.__table__.alias("history_seed")
    seed = select(
        seed_root.c.id.label("root_id"), seed_root.c.id.label("node_id"),
        literal(0).label("depth"),
        (literal(",") + cast(seed_root.c.id, String) + literal(",")).label("visited"),
        cast(literal(None), Integer).label("proof_id"),
    ).where(seed_root.c.data["cross_session_suppressed"].as_boolean().is_(True)).cte(
        "vacancy_history_walk", recursive=True
    )
    current = Vacancy.__table__.alias("history_current")
    target = Vacancy.__table__.alias("history_target")
    historical_id = func.coalesce(
        cast(func.json_extract(current.c.data, "$.historical_vacancy_id"), String),
        cast(func.json_extract(current.c.data, "$.historical_source_vacancy_id"), String),
    )
    target_int = cast(historical_id, Integer)
    current_application = exists(select(Application.id).where(
        Application.vacancy_id == current.c.id,
        Application.status.in_(("submitted", "already_applied")),
    ))
    current_cv = func.json_extract(current.c.data, "$.submission_progress.cv_confirmed") == 1
    current_letter_confirmed = func.json_extract(current.c.data, "$.submission_progress.cover_letter_confirmed") == 1
    current_letter_pending = func.json_extract(current.c.data, "$.submission_progress.cover_letter_pending") == 1
    current_proof = current_application | current.c.state.in_(["SUBMITTED", "REPORTED"]) | (current_cv & current_letter_confirmed)
    current_partial = (current.c.state == "PARTIAL") | (current_cv & current_letter_pending)
    next_row = select(
        seed.c.root_id, target.c.id.label("node_id"), (seed.c.depth + 1).label("depth"),
        (seed.c.visited + cast(target.c.id, String) + literal(",")).label("visited"),
        func.coalesce(seed.c.proof_id, case((current_proof | current_partial, current.c.id), else_=None)).label("proof_id"),
    ).select_from(
        seed.join(current, current.c.id == seed.c.node_id).join(target, target.c.id == target_int)
    ).where(
        seed.c.depth < 32,
        case(
            (func.json_extract(current.c.data, "$.historical_vacancy_id").is_not(None), func.json_type(current.c.data, "$.historical_vacancy_id")),
            else_=func.json_type(current.c.data, "$.historical_source_vacancy_id"),
        ) == "integer",
        target.c.id < current.c.id,
        target.c.source == current.c.source,
        target.c.external_id == current.c.external_id,
        target.c.external_id.is_not(None),
        func.instr(seed.c.visited, literal(",") + cast(target.c.id, String) + literal(",")) == 0,
    )
    walk = seed.union_all(next_row)
    deepest = select(walk.c.root_id, func.max(walk.c.depth).label("depth")).group_by(walk.c.root_id).subquery()
    leaf = walk.join(deepest, (walk.c.root_id == deepest.c.root_id) & (walk.c.depth == deepest.c.depth))
    origin = Vacancy.__table__.alias("history_origin")
    root = Vacancy.__table__.alias("history_root")
    origin_data = origin.c.data
    root_data = root.c.data
    leaf_next = func.coalesce(
        func.json_extract(origin_data, "$.historical_vacancy_id"),
        func.json_extract(origin_data, "$.historical_source_vacancy_id"),
    )
    origin_application = exists(select(Application.id).where(
        Application.vacancy_id == origin.c.id,
        Application.status.in_(("submitted", "already_applied")),
    ))
    origin_cv = func.json_extract(origin_data, "$.submission_progress.cv_confirmed") == 1
    origin_letter_confirmed = func.json_extract(origin_data, "$.submission_progress.cover_letter_confirmed") == 1
    origin_letter_pending = func.json_extract(origin_data, "$.submission_progress.cover_letter_pending") == 1
    origin_proof = origin_application | origin.c.state.in_(["SUBMITTED", "REPORTED"]) | (origin_cv & origin_letter_confirmed)
    origin_partial = (origin.c.state == "PARTIAL") | (origin_cv & origin_letter_pending)
    invalid = case(
        (origin_proof | origin_partial, False),
        (walk.c.proof_id.is_not(None), False),
        (leaf_next.is_not(None), True),
        else_=False,
    )
    root_application = exists(select(Application.id).where(
        Application.vacancy_id == root.c.id,
        Application.status.in_(("submitted", "already_applied")),
    ))
    root_cv = func.json_extract(root_data, "$.submission_progress.cv_confirmed") == 1
    root_letter_confirmed = func.json_extract(root_data, "$.submission_progress.cover_letter_confirmed") == 1
    root_letter_pending = func.json_extract(root_data, "$.submission_progress.cover_letter_pending") == 1
    origin_code = func.coalesce(origin.c.error_code, func.json_extract(origin_data, "$.error_code"))
    outcome = case(
        (invalid, "unconfirmed"),
        ((root_application | root.c.state.in_(["SUBMITTED", "REPORTED"]) | (root_cv & root_letter_confirmed)), "confirmed"),
        ((root_cv & root_letter_pending) | (root.c.state == "PARTIAL"), "partial"),
        (origin_proof, "confirmed"),
        (origin_partial, "partial"),
        (origin.c.state == "ALREADY_APPLIED", "already_applied"),
        ((origin.c.state == "PARTIAL") | ((func.json_extract(origin_data, "$.submission_progress.cv_confirmed") == 1) & (func.json_extract(origin_data, "$.submission_progress.cover_letter_pending") == 1)), "partial"),
        ((origin.c.state.in_(["SUBMISSION_UNCONFIRMED", "UNCONFIRMED", "SUBMITTING"]) | (origin_code == "SUBMISSION_UNCONFIRMED") | (func.json_extract(origin_data, "$.submission_attempted") == 1) | (func.json_extract(origin_data, "$.recovery_unresolved") == 1) | (func.json_extract(origin_data, "$.partial_recovery_blocked") == 1) | (func.json_extract(origin_data, "$.submission_reconciliation_attempts") > 0)), "unconfirmed"),
        else_="unconfirmed",
    )
    effective = case(
        (invalid, "SUBMISSION_UNCONFIRMED"),
        (outcome == "partial", "PARTIAL"),
        (outcome == "unconfirmed", "SUBMISSION_UNCONFIRMED"),
        (outcome == "already_applied", "ALREADY_APPLIED"),
        ((outcome == "confirmed") & root.c.state.in_(["SUBMISSION_UNCONFIRMED", "UNCONFIRMED", "SUBMITTING", "EXTRACTED", "EVALUATING"]), "SUBMITTED"),
        else_=root.c.state,
    )
    return select(
        walk.c.root_id.label("vacancy_id"), origin.c.id.label("origin_id"),
        origin.c.session_id.label("origin_session_id"), invalid.label("invalid_history"),
        outcome.label("history_outcome"), effective.label("effective_state"),
    ).select_from(
        leaf.join(origin, origin.c.id == func.coalesce(walk.c.proof_id, walk.c.node_id)).join(root, root.c.id == walk.c.root_id)
    ).cte("vacancy_history_projection")


def _row_payload(row, include_data: bool = False) -> dict:
    effective_state = getattr(row, "effective_state", row.state)
    state = _public_status(effective_state)
    history_outcome = getattr(row, "history_outcome", None)
    result = {
        "id": row.id, "session_id": row.session_id, "title": row.title, "company": row.company,
        "url": row.url, "state": state,
        "status_group": next((name for name, states in VACANCY_STATUS_GROUPS.items() if effective_state in states), "ERROR"),
        "source": row.source or "", "site": row.site or "", "status_changed_at": _status_time(row.status_changed_at, row.site),
        "evaluation": _evaluation_summary(row, getattr(row, "evaluation_data", None) if include_data else None),
    }
    if getattr(row, "history_origin_id", None) is not None or getattr(row, "history_invalid", False):
        result["analysis_status"] = "not_evaluated_history"
        if getattr(row, "history_origin_id", None) is not None:
            result["history_context"] = {
                "source_vacancy_id": row.history_origin_id,
                "source_session_id": row.history_origin_session_id,
                "outcome": history_outcome or "unconfirmed",
                "reason_code": "cross_session_suppressed",
            }
        else:
            result["analysis_status"] = "not_evaluated_history"
            result["history_context"] = {
                "source_vacancy_id": None, "source_session_id": None,
                "outcome": "unconfirmed", "reason_code": "invalid_history",
            }
    elif row.evaluation_id is not None:
        result["analysis_status"] = "scored"
    elif effective_state in {"EXTRACTED", "EVALUATING", "READY_TO_SUBMIT", "READY_TO_REPORT", "SUBMITTING"}:
        result["analysis_status"] = "pending"
    else:
        result["analysis_status"] = "not_evaluated"
    if effective_state == "SUBMISSION_UNCONFIRMED":
        result["stored_state"] = row.state
    if state == "ERROR" or effective_state == "SUBMISSION_UNCONFIRMED":
        result["error_code"], result["error_message"] = _error(row)
    if include_data:
        result["data"] = row.vacancy_data or {}
    return result


def _score_limits(**values) -> dict[str, tuple[float | None, float | None]]:
    return {key: (values.get(f"{key}_min"), values.get(f"{key}_max")) for key in VACANCY_SCORE_KEYS}


def _where(*, search=None, state=None, status_group=None, site=None, status_date_from=None, status_date_to=None,
           status_time_from=None, status_time_before=None,
           total_score_min=None, total_score_max=None, score_limits=None, state_column=None):
    clauses = []
    if search and search.strip():
        needle = search.strip().casefold().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        clauses.append(or_(cast(Vacancy.id, String).like(f"%{needle}%", escape="\\"), Vacancy.search_text.like(f"%{needle}%", escape="\\")))
    accepted = set(VACANCY_STATUS_GROUPS[status_group]) if status_group else None
    if state:
        requested = set(VACANCY_STATUS_GROUPS.get(state, {state}))
        accepted = requested if accepted is None else accepted & requested
    if accepted is not None:
        clauses.append((state_column if state_column is not None else Vacancy.state).in_(accepted))
    if site == "__legacy__":
        clauses.append(Vacancy.site == "")
    elif site is not None:
        clauses.append(Vacancy.site == site)
    if status_date_from:
        clauses.append(Vacancy.status_changed_at >= datetime.combine(status_date_from, time.min))
    if status_date_to:
        clauses.append(Vacancy.status_changed_at <= datetime.combine(status_date_to, time.max))
    if status_time_from is not None:
        clauses.append(Vacancy.status_changed_at >= status_time_from)
    if status_time_before is not None:
        clauses.append(Vacancy.status_changed_at < status_time_before)
    if total_score_min is not None:
        clauses.append(Evaluation.total_score >= total_score_min)
    if total_score_max is not None:
        clauses.append(Evaluation.total_score <= total_score_max)
    for key, (minimum, maximum) in (score_limits or {}).items():
        column = VACANCY_SCORE_COLUMNS[key]
        if minimum is not None:
            clauses.append(column >= minimum)
        if maximum is not None:
            clauses.append(column <= maximum)
    return clauses


def _select(include_data: bool = False, history=None):
    columns = [
        Vacancy.id, Vacancy.session_id, Vacancy.source, Vacancy.site, Vacancy.external_id, Vacancy.url,
        Vacancy.title, Vacancy.company, Vacancy.state, Vacancy.status_changed_at, Vacancy.error_code,
        Vacancy.error_message, Evaluation.id.label("evaluation_id"), Evaluation.total_score,
        Evaluation.decision, Evaluation.confidence, Evaluation.category,
        *(getattr(Evaluation, key) for key in VACANCY_SCORE_KEYS),
    ]
    if include_data:
        columns.extend([Vacancy.data.label("vacancy_data"), Evaluation.data.label("evaluation_data")])
    history = history if history is not None else _history_projection()
    columns.extend([
        case((history.c.vacancy_id.is_not(None), history.c.effective_state), else_=Vacancy.state).label("effective_state"),
        history.c.origin_id.label("history_origin_id"),
        history.c.origin_session_id.label("history_origin_session_id"),
        history.c.invalid_history.label("history_invalid"),
        history.c.history_outcome.label("history_outcome"),
    ])
    return select(*columns).outerjoin(history, history.c.vacancy_id == Vacancy.id).outerjoin(Evaluation, Evaluation.vacancy_id == Vacancy.id)


def _ordered(stmt, sort: VacancySort, sort_dir: VacancySortDirection, *, tie_column=None, state_column=None):
    if sort == "id":
        column = Vacancy.id
    elif sort == "title":
        column = Vacancy.title_sort
    elif sort == "site":
        column = Vacancy.site_sort
    elif sort == "state":
        # UNCONFIRMED is retained in storage for old rows but is exposed as
        # ERROR. Sort on the same value the list API returns.
        raw_state = state_column if state_column is not None else Vacancy.state
        column = case((raw_state == "UNCONFIRMED", "ERROR"), else_=raw_state)
    elif sort == "date":
        column = Vacancy.status_changed_at
    else:
        column = VACANCY_SCORE_COLUMNS[sort]
    direction = column.asc() if sort_dir == "asc" else column.desc()
    # ``NULLS LAST`` is understood by SQLite and PostgreSQL and keeps the
    # leading sort key indexable.  The previous CASE expression forced a
    # temporary B-tree for every list, even when a covering projection index
    # was available.  The vacancy id is the stable, direction-independent
    # tie-breaker used by both the list and export paths.
    return stmt.order_by(direction.nulls_last(), (tie_column or Vacancy.id).asc())


def _statement(*, include_data=False, sort="date", sort_dir="desc", **filters):
    history = _history_projection()
    state_column = case((history.c.vacancy_id.is_not(None), history.c.effective_state), else_=Vacancy.state)
    return _ordered(_select(include_data, history=history).where(*_where(**filters, state_column=state_column)), sort, sort_dir, state_column=state_column)


def _query(db: Session, *, limit=None, offset=0, sort="date", sort_dir="desc", include_data=False, **filters):
    stmt = _statement(include_data=include_data, sort=sort, sort_dir=sort_dir, **filters)
    if limit is None:
        return db.execute(stmt).mappings().all()

    # Find the page using the narrow scalar projection first.  In particular,
    # this prevents a detail request (include_data=True) from carrying large
    # JSON values through the sort before LIMIT/OFFSET has reduced the page.
    # The outer query only joins those rows back to fetch the response shape.
    # A score range excludes NULL scores, so the evaluation covering index
    # can provide both the score and its vacancy-id tie-breaker directly.
    # Without a range we retain Vacancy.id as the tie-breaker so unscored
    # vacancies remain deterministic after the NULL partition.
    score_range = filters.get("total_score_min") is not None or filters.get("total_score_max") is not None
    if sort != "total_score":
        score_range = any(value is not None for value in filters.get("score_limits", {}).get(sort, (None, None)))
    page_ids = None
    # An outer join cannot use the evaluation projection index for an
    # unfiltered score sort: SQLite starts at vacancies and materialises a
    # full TEMP B-tree to put NULLs last.  Split the order into its natural
    # partitions instead.  The non-NULL partition is index-backed, while the
    # NULL partition is already deterministically ordered by the vacancy PK.
    if sort in VACANCY_SCORE_COLUMNS and not score_range:
        score_column = VACANCY_SCORE_COLUMNS[sort]
        history = _history_projection()
        effective_state = case((history.c.vacancy_id.is_not(None), history.c.effective_state), else_=Vacancy.state)
        clauses = _where(**filters, state_column=effective_state)
        non_null = [*clauses, score_column.is_not(None)]
        vacancy_filter = any(filters[key] is not None for key in (
            "search", "state", "status_group", "site", "status_date_from", "status_date_to",
            "status_time_from", "status_time_before",
        ))
        count_stmt = select(func.count()).select_from(Evaluation)
        if vacancy_filter:
            count_stmt = count_stmt.join(Vacancy, Evaluation.vacancy_id == Vacancy.id).outerjoin(history, history.c.vacancy_id == Vacancy.id)
        non_null_count = db.scalar(count_stmt.where(*non_null)) or 0
        remaining = limit
        if offset < non_null_count:
            non_null_limit = min(limit, non_null_count - offset)
            score_stmt = select(Evaluation.vacancy_id)
            if vacancy_filter:
                score_stmt = score_stmt.join(Vacancy, Evaluation.vacancy_id == Vacancy.id).outerjoin(history, history.c.vacancy_id == Vacancy.id)
            direction = score_column.asc() if sort_dir == "asc" else score_column.desc()
            score_stmt = score_stmt.where(*non_null).order_by(direction, Evaluation.vacancy_id.asc())
            page_ids = db.execute(score_stmt.limit(non_null_limit).offset(offset)).scalars().all()
            remaining -= len(page_ids)
        if remaining:
            null_clauses = [*clauses, score_column.is_(None)]
            null_stmt = select(Vacancy.id).select_from(Vacancy).outerjoin(
                Evaluation, Evaluation.vacancy_id == Vacancy.id
            ).outerjoin(history, history.c.vacancy_id == Vacancy.id).where(*null_clauses).order_by(Vacancy.id.asc())
            null_ids = db.execute(null_stmt.limit(remaining).offset(max(0, offset - non_null_count))).scalars().all()
            page_ids = (page_ids or []) + null_ids

    if page_ids is None:
        tie_column = Evaluation.vacancy_id if sort in VACANCY_SCORE_COLUMNS and score_range else None
        history = _history_projection()
        effective_state = case((history.c.vacancy_id.is_not(None), history.c.effective_state), else_=Vacancy.state)
        page_ids_stmt = _ordered(
            select(Vacancy.id).select_from(Vacancy).outerjoin(history, history.c.vacancy_id == Vacancy.id).outerjoin(Evaluation, Evaluation.vacancy_id == Vacancy.id)
            .where(*_where(**filters, state_column=effective_state)),
            sort,
            sort_dir,
            tie_column=tie_column,
            state_column=effective_state,
        ).limit(limit).offset(offset)
        page_ids = db.execute(page_ids_stmt).scalars().all()
    if not page_ids:
        return []

    # Fetch the (possibly JSON-bearing) projection only for the selected
    # primary keys.  Reordering in Python preserves the exact page order while
    # avoiding a second SQL sort after the database has already selected the
    # page.  The old join-back ORDER BY created a TEMP B-tree even when the
    # narrow page query was fully index-backed.
    rows = db.execute(_select(include_data).where(Vacancy.id.in_(page_ids))).mappings().all()
    by_id = {row.id: row for row in rows}
    return [by_id[vacancy_id] for vacancy_id in page_ids if vacancy_id in by_id]


def _iter_rows(bind, *, sort, sort_dir, filters) -> Iterator:
    # StreamingResponse may outlive the request dependency.  Open a dedicated
    # session inside the generator so the cursor remains valid until the
    # iterator is exhausted or cancelled, then close both deterministically.
    stream_factory = sessionmaker(bind=bind, autoflush=False, expire_on_commit=False)
    with stream_factory() as stream_db:
        result = None
        try:
            result = stream_db.execute(_statement(sort=sort, sort_dir=sort_dir, **filters).execution_options(stream_results=True, yield_per=_EXPORT_BATCH_SIZE)).mappings()
            while chunk := result.fetchmany(_EXPORT_BATCH_SIZE):
                yield from chunk
        finally:
            if result is not None:
                result.close()


def _filters(*, search, state, status_group, site, status_date_from=None, status_date_to=None,
             status_time_from=None, status_time_before=None, total_score_min=None,
             total_score_max=None, tasks_min=None, tasks_max=None, skills_min=None, skills_max=None,
             experience_depth_min=None, experience_depth_max=None, role_match_min=None, role_match_max=None,
             industry_min=None, industry_max=None, special_requirements_min=None, special_requirements_max=None):
    return {
        "search": search, "state": state, "status_group": status_group, "site": site,
        "status_date_from": status_date_from, "status_date_to": status_date_to,
        "status_time_from": status_time_from, "status_time_before": status_time_before,
        "total_score_min": total_score_min, "total_score_max": total_score_max,
        "score_limits": _score_limits(
            tasks_min=tasks_min, tasks_max=tasks_max, skills_min=skills_min, skills_max=skills_max,
            experience_depth_min=experience_depth_min, experience_depth_max=experience_depth_max,
            role_match_min=role_match_min, role_match_max=role_match_max, industry_min=industry_min,
            industry_max=industry_max, special_requirements_min=special_requirements_min,
            special_requirements_max=special_requirements_max,
        ),
    }


def _filter_params(search, state, status_group, site, status_date_from, status_date_to, total_score_min, total_score_max,
                   tasks_min, tasks_max, skills_min, skills_max, experience_depth_min, experience_depth_max,
                   role_match_min, role_match_max, industry_min, industry_max, special_requirements_min,
                   special_requirements_max, status_time_from=None, status_time_before=None):
    if (status_time_from is not None or status_time_before is not None) and (
        status_date_from is not None or status_date_to is not None
    ):
        raise HTTPException(422, "Не смешивайте два способа задания периода")
    normalized_time_from = _utc_naive_bound(status_time_from, "status_time_from")
    normalized_time_before = _utc_naive_bound(status_time_before, "status_time_before")
    if (
        normalized_time_from is not None
        and normalized_time_before is not None
        and normalized_time_from >= normalized_time_before
    ):
        raise HTTPException(422, "Начало периода должно быть раньше его окончания")
    return _filters(search=search, state=state, status_group=status_group, site=site, status_date_from=status_date_from,
                    status_time_from=normalized_time_from, status_time_before=normalized_time_before,
                    status_date_to=status_date_to, total_score_min=total_score_min, total_score_max=total_score_max,
                    tasks_min=tasks_min, tasks_max=tasks_max, skills_min=skills_min, skills_max=skills_max,
                    experience_depth_min=experience_depth_min, experience_depth_max=experience_depth_max,
                    role_match_min=role_match_min, role_match_max=role_match_max, industry_min=industry_min,
                    industry_max=industry_max, special_requirements_min=special_requirements_min,
                    special_requirements_max=special_requirements_max)


def _utc_naive_bound(value: datetime | None, name: str) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None or value.utcoffset() is None:
        raise HTTPException(422, "Для границы периода требуется часовой пояс")
    return value.astimezone(timezone.utc).replace(tzinfo=None)


@router.get("/vacancies")
def vacancies(limit: int = Query(30, ge=1, le=100), offset: int = Query(0, ge=0), include_data: bool = Query(False),
              search: str | None = None, state: str | None = None, status_group: VacancyStatusGroup | None = None,
              site: str | None = None, status_date_from: date | None = None, status_date_to: date | None = None,
              status_time_from: datetime | None = None, status_time_before: datetime | None = None,
              total_score_min: float | None = Query(None, ge=0, le=100), total_score_max: float | None = Query(None, ge=0, le=100),
              sort: VacancySort = "date", sort_dir: VacancySortDirection = "desc", tasks_min: float | None = None,
              tasks_max: float | None = None, skills_min: float | None = None, skills_max: float | None = None,
              experience_depth_min: float | None = None, experience_depth_max: float | None = None,
              role_match_min: float | None = None, role_match_max: float | None = None, industry_min: float | None = None,
              industry_max: float | None = None, special_requirements_min: float | None = None,
              special_requirements_max: float | None = None, db: Session = Depends(get_db)):
    filters = _filter_params(search, state, status_group, site, status_date_from, status_date_to, total_score_min,
                             total_score_max, tasks_min, tasks_max, skills_min, skills_max, experience_depth_min,
                             experience_depth_max, role_match_min, role_match_max, industry_min, industry_max,
                             special_requirements_min, special_requirements_max,
                             status_time_from, status_time_before)
    history = _history_projection()
    effective_state = case((history.c.vacancy_id.is_not(None), history.c.effective_state), else_=Vacancy.state)
    where = _where(**filters, state_column=effective_state)
    score_filter = (
        filters["total_score_min"] is not None or filters["total_score_max"] is not None
        or any(value is not None for limits in filters["score_limits"].values() for value in limits)
    )
    # A score-only count can use the covering evaluation index directly.  The
    # former unconditional vacancy LEFT JOIN visited tens of thousands of
    # vacancy rows (and their primary-key lookups) just to count evaluations.
    vacancy_filter = any(filters[key] is not None for key in (
        "search", "state", "status_group", "site", "status_date_from", "status_date_to",
        "status_time_from", "status_time_before",
    ))
    if score_filter and not vacancy_filter:
        count_stmt = select(func.count()).select_from(Evaluation).where(*where)
    else:
        count_stmt = select(func.count()).select_from(Vacancy).outerjoin(history, history.c.vacancy_id == Vacancy.id)
        if score_filter:
            count_stmt = count_stmt.join(Evaluation, Evaluation.vacancy_id == Vacancy.id)
        count_stmt = count_stmt.where(*where)
    total = db.scalar(count_stmt) or 0
    rows = _query(db, limit=limit, offset=offset, sort=sort, sort_dir=sort_dir, include_data=include_data, **filters)
    return {"items": [_row_payload(row, include_data) for row in rows], "total": total, "limit": limit,
            "offset": offset, "has_more": offset + len(rows) < total}


def _export_values(row) -> list[object]:
    def number(value):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    return [row.id, row.title, row.company or "", row.site or "", _export_status_time(row.status_changed_at, row.site),
            number(row.total_score) if row.total_score is not None else "", *[number(getattr(row, key)) if getattr(row, key) is not None else "" for key in VACANCY_SCORE_KEYS]]


def _csv_body(rows: Iterator) -> Iterator[bytes]:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(VACANCY_EXPORT_HEADERS)
    yield output.getvalue().encode("utf-8-sig")
    try:
        for row in rows:
            output.seek(0); output.truncate(0); writer.writerow(_export_values(row))
            yield output.getvalue().encode("utf-8")
    finally:
        close = getattr(rows, "close", None)
        if close is not None:
            close()


def _xml_body(rows: Iterator) -> Iterator[bytes]:
    yield b'<?xml version="1.0" encoding="utf-8"?>\n<vacancies>'
    try:
        for row in rows:
            item = ET.Element("vacancy")
            for header, value in zip(VACANCY_EXPORT_HEADERS, _export_values(row), strict=True):
                field = ET.SubElement(item, "field", name=header); field.text = "" if value is None else str(value)
            yield ET.tostring(item, encoding="utf-8")
        yield b"</vacancies>"
    finally:
        close = getattr(rows, "close", None)
        if close is not None:
            close()


def _xlsx_body(rows: Iterator) -> Iterator[bytes]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        path = Path(tmp.name)
    book = None
    try:
        book = Workbook(write_only=True); sheet = book.create_sheet("Вакансии"); sheet.append(VACANCY_EXPORT_HEADERS)
        for row in rows:
            sheet.append(_export_values(row))
        book.save(path)
        with path.open("rb") as stream:
            while chunk := stream.read(64 * 1024):
                yield chunk
    finally:
        if book is not None:
            book.close()
        path.unlink(missing_ok=True)
        close = getattr(rows, "close", None)
        if close is not None:
            close()


@router.get("/vacancies/export")
def export_vacancies(format: VacancyExportFormat = "csv", search: str | None = None, state: str | None = None,
                     status_group: VacancyStatusGroup | None = None, site: str | None = None,
                     status_date_from: date | None = None, status_date_to: date | None = None,
                     status_time_from: datetime | None = None, status_time_before: datetime | None = None,
                     total_score_min: float | None = Query(None, ge=0, le=100), total_score_max: float | None = Query(None, ge=0, le=100),
                     sort: VacancySort = "date", sort_dir: VacancySortDirection = "desc", tasks_min: float | None = None,
                     tasks_max: float | None = None, skills_min: float | None = None, skills_max: float | None = None,
                     experience_depth_min: float | None = None, experience_depth_max: float | None = None,
                     role_match_min: float | None = None, role_match_max: float | None = None, industry_min: float | None = None,
                     industry_max: float | None = None, special_requirements_min: float | None = None,
                     special_requirements_max: float | None = None, db: Session = Depends(get_db)):
    filters = _filter_params(search, state, status_group, site, status_date_from, status_date_to, total_score_min,
                             total_score_max, tasks_min, tasks_max, skills_min, skills_max, experience_depth_min,
                             experience_depth_max, role_match_min, role_match_max, industry_min, industry_max,
                             special_requirements_min, special_requirements_max,
                             status_time_from, status_time_before)
    rows = _iter_rows(db.get_bind(), sort=sort, sort_dir=sort_dir, filters=filters)
    headers = {"Content-Disposition": f'attachment; filename="vacancies.{format}"'}
    if format == "xlsx":
        return StreamingResponse(_xlsx_body(rows), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)
    if format == "xml":
        return StreamingResponse(_xml_body(rows), media_type="application/xml", headers=headers)
    return StreamingResponse(_csv_body(rows), media_type="text/csv; charset=utf-8", headers=headers)


@router.get("/vacancies/{vacancy_id}")
def vacancy_detail(vacancy_id: int, db: Session = Depends(get_db)):
    row = db.execute(_select(include_data=True).where(Vacancy.id == vacancy_id).limit(1)).mappings().first()
    if row is None:
        raise HTTPException(404, "Вакансия не найдена")
    return Response(
        content=json.dumps(_row_payload(row, True), ensure_ascii=False, default=str),
        media_type="application/json",
        headers={"Cache-Control": "no-store"},
    )
